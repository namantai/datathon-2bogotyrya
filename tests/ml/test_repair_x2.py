"""Synthetic workbook-only repair tests; never access organizer data."""
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import pytest

from ml.repair_x2 import EXPECTED, RepairError, repair


def fixture(path: Path, row=None, placeholder=False):
    wb = Workbook()
    ws = wb.active
    ws.title = 'TRAIN_27'
    headers = list(reversed(EXPECTED)) + ['observed_action_ru', 'formula']
    ws.append(headers)
    ws.append([{'video_id': 'SYNTHETIC-1', 'filename': 'SYNTHETIC-1.mov',
                'label': 0, 'formula': '=1+2'}.get(h) for h in headers])
    if row is not None:
        ws.append([row.get(h) for h in headers])
    if placeholder:
        ws.cell(3, 1).fill = PatternFill('solid', fgColor='ABCDEF')
    wb.create_sheet('Other')['A1'] = '=2+3'
    wb.save(path)
    wb.close()
    return headers


def test_append_preserves_zip_parts_and_second_run_is_byte_identical(tmp_path):
    p = tmp_path / 'synthetic.xlsx'
    fixture(p)
    original = p.read_bytes()
    result = repair(p)
    assert result['backup_verified'] and result['other_parts_preserved']
    backups = list(tmp_path.glob('*.backup-*.xlsx'))
    assert len(backups) == 1 and backups[0].read_bytes() == original
    with ZipFile(backups[0]) as a, ZipFile(p) as b:
        assert a.namelist() == b.namelist()
        assert [n for n in a.namelist() if a.read(n) != b.read(n)] == ['xl/worksheets/sheet1.xml']
    changed = p.read_bytes()
    assert not repair(p)['changed']
    assert p.read_bytes() == changed
    assert len(list(tmp_path.glob('*.backup-*.xlsx'))) == 1


def test_reuse_empty_styled_placeholder(tmp_path):
    p = tmp_path / 'synthetic.xlsx'
    headers = fixture(p, placeholder=True)
    repair(p)
    wb = load_workbook(p)
    ws = wb['TRAIN_27']
    assert ws.max_row == 3
    assert ws.cell(3, 1).fill.fgColor.rgb == '00ABCDEF'
    for key, value in EXPECTED.items():
        assert ws.cell(3, headers.index(key) + 1).value == value
    assert ws.cell(3, headers.index('observed_action_ru') + 1).value is None
    wb.close()


def test_preserve_nonempty_unidentified_row(tmp_path):
    p = tmp_path / 'synthetic.xlsx'
    fixture(p, {'observed_action_ru': 'SYNTHETIC PLACEHOLDER NOTE'})
    repair(p)
    wb = load_workbook(p)
    assert wb['TRAIN_27'].max_row == 4
    assert 'SYNTHETIC PLACEHOLDER NOTE' in [c.value for c in wb['TRAIN_27'][3]]
    wb.close()


def test_normalized_existing_id_and_timestamp_do_not_duplicate(tmp_path):
    p = tmp_path / 'synthetic.xlsx'
    fixture(p, {'video_id': ' х-2.MOV ', 'timecode_sec': '00:10', 'label': '1'})
    assert repair(p)['existing_row_updated']
    wb = load_workbook(p)
    assert wb['TRAIN_27'].max_row == 3
    wb.close()
    assert not repair(p)['changed']


@pytest.mark.parametrize('extra', [
    {'label': 0}, {'timecode_sec': 11}, {'event_type': 'unverified'},
    {'filename': 'DIFFERENT.mov'}, {'observed_action_ru': 'SYNTHETIC NOTE'},
    {'label': '=1'},
])
def test_conflicts_leave_source_untouched_without_backup(tmp_path, extra):
    p = tmp_path / 'synthetic.xlsx'
    fixture(p, {'video_id': 'X-2', **extra})
    original = p.read_bytes()
    with pytest.raises(RepairError, match='conflicting_x2'):
        repair(p)
    assert p.read_bytes() == original
    assert not list(tmp_path.glob('*.backup-*.xlsx'))


def test_duplicate_rows_stop_before_any_write(tmp_path):
    p = tmp_path / 'synthetic.xlsx'
    headers = fixture(p, {'video_id': 'X-2'})
    wb = load_workbook(p)
    wb['TRAIN_27'].append([{'filename': 'Х-2.MOV'}.get(h) for h in headers])
    wb.save(p)
    wb.close()
    original = p.read_bytes()
    with pytest.raises(RepairError, match='conflicting_x2_duplicate_rows'):
        repair(p)
    assert p.read_bytes() == original
    assert not list(tmp_path.glob('*.backup-*.xlsx'))
