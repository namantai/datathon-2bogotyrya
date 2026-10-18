"""Local, idempotent X-2 repair; preserve every other XLSX ZIP member verbatim.

No private cell values are logged. Nonblank conflicting annotations stop the repair.
The original is backed up before atomic replacement; no backup is overwritten.
"""
from __future__ import annotations

import argparse
from copy import copy
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import posixpath
import re
import stat
import uuid
import warnings
from xml.dom import minidom
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from ml.dataset import normalize_video_id, parse_label, parse_timestamp
from ml.local_io import private_path

EXPECTED = {
    'video_id': 'X-2', 'filename': 'X-2.mov', 'label': 1,
    'official_label': 'нарушение', 'timecode_sec': 10,
    'event_group': 'unknown', 'event_type': 'unknown_violation',
    'severity': 'unknown', 'qa_status': 'needs_review', 'split': 'train',
    'annotation_source': 'official_label',
}
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


class RepairError(ValueError):
    """Messages are fixed technical codes, never cell contents."""


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def compatible(key, value):
    if blank(value):
        return True
    try:
        if key in ('video_id', 'filename'):
            return normalize_video_id(value) == 'X-2'
        if key == 'label':
            return parse_label(value, {'1'}, {'0'}) == 1
        if key == 'timecode_sec':
            return parse_timestamp(value) == 10
        return isinstance(value, str) and value.strip() == EXPECTED[key]
    except (ValueError, TypeError):
        return False


def plan(original):
    with warnings.catch_warnings():
        warnings.simplefilter('error')
        wb = load_workbook(BytesIO(original), data_only=False, read_only=True)
        try:
            if 'TRAIN_27' not in wb.sheetnames:
                raise RepairError('missing_sheet')
            ws = wb['TRAIN_27']
            rows = list(ws.values)
            candidates = [(i, row) for i, row in enumerate(rows, 1)
                          if all(row.count(key) == 1 for key in EXPECTED)]
            if len(candidates) != 1:
                raise RepairError('ambiguous_headers')
            header, names = candidates[0]
            columns = {key: names.index(key) + 1 for key in EXPECTED}
            matches, empty = [], []
            for number, row in enumerate(rows[header:], header + 1):
                if all(blank(v) for v in row):
                    empty.append(number)
                for key in ('video_id', 'filename'):
                    value = row[columns[key] - 1]
                    try:
                        match = not blank(value) and normalize_video_id(value) == 'X-2'
                    except ValueError:
                        match = False
                    if match:
                        matches.append(number)
                        break
            if len(matches) > 1:
                raise RepairError('conflicting_x2_duplicate_rows')
            target = matches[0] if matches else empty[0] if empty else len(rows) + 1
            changes = {}
            if matches:
                row = rows[target - 1]
                if any(not compatible(k, row[c - 1]) for k, c in columns.items()):
                    raise RepairError('conflicting_x2_fields')
                if any(not blank(v) for c, v in enumerate(row, 1) if c not in columns.values()):
                    raise RepairError('conflicting_x2_unverified_fields')
            for key, column in columns.items():
                value = rows[target - 1][column - 1] if target <= len(rows) else None
                if value != EXPECTED[key] or isinstance(value, bool):
                    changes[f'{get_column_letter(column)}{target}'] = EXPECTED[key]
            return target, changes, bool(matches)
        finally:
            wb.close()


def patch_archive(original, target, changes):
    """Edit only the target row and (when needed) worksheet dimension bytes."""
    with ZipFile(BytesIO(original)) as source:
        if len(source.namelist()) != len(set(source.namelist())):
            raise RepairError('duplicate_zip_members')
        workbook = ET.fromstring(source.read('xl/workbook.xml'))
        sheet = next(s for s in workbook.find(f'{{{NS}}}sheets') if s.attrib['name'] == 'TRAIN_27')
        rid = sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        relations = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        destination = next(r.attrib['Target'] for r in relations if r.attrib['Id'] == rid)
        member = destination.lstrip('/') if destination.startswith('/') else posixpath.normpath('xl/' + destination)
        raw = source.read(member)
        doc = minidom.parseString(raw)
        data = doc.getElementsByTagNameNS(NS, 'sheetData')[0]
        # Fail closed on unusual namespace prefixes rather than rewrite other XML.
        if data.tagName != 'sheetData':
            raise RepairError('unsupported_sheet_prefix')
        old = next((n for n in data.childNodes if n.nodeType == n.ELEMENT_NODE
                    and n.getAttribute('r') == str(target)), None)
        node = old.cloneNode(True) if old else doc.createElement('row')
        node.setAttribute('r', str(target))
        for reference, value in changes.items():
            cell = next((c for c in node.childNodes if c.nodeType == c.ELEMENT_NODE
                         and c.getAttribute('r') == reference), None)
            if cell is None:
                cell = doc.createElement('c')
                cell.setAttribute('r', reference)
                node.appendChild(cell)
            for child in list(cell.childNodes):
                cell.removeChild(child)
            cell.setAttribute('t', 'n' if isinstance(value, int) else 'inlineStr')
            payload = doc.createElement('v' if isinstance(value, int) else 'is')
            if isinstance(value, int):
                payload.appendChild(doc.createTextNode(str(value)))
            else:
                text = doc.createElement('t')
                text.appendChild(doc.createTextNode(value))
                payload.appendChild(text)
            cell.appendChild(payload)
        # Keep cells in numeric column order, including columns beyond Z.
        cells = [c for c in node.childNodes if c.nodeType == c.ELEMENT_NODE and c.tagName == 'c']
        for cell in sorted(cells, key=lambda c: range_boundaries(c.getAttribute('r'))[0]):
            node.appendChild(cell)
        replacement = node.toxml(encoding='utf-8')
        if old:
            pattern = rb'<row\b(?=[^>]*\br="' + str(target).encode() + rb'")[^>]*(?:/>|>.*?</row>)'
            updated, count = re.subn(pattern, lambda _: replacement, raw, flags=re.S)
            if count != 1:
                raise RepairError('ambiguous_row_xml')
        else:
            if raw.count(b'</sheetData>') != 1:
                raise RepairError('unsupported_sheet_data')
            updated = raw.replace(b'</sheetData>', replacement + b'</sheetData>', 1)
        dimensions = doc.getElementsByTagNameNS(NS, 'dimension')
        if dimensions:
            dim = dimensions[0]
            left, top, right, bottom = range_boundaries(dim.getAttribute('ref'))
            new_ref = f'{get_column_letter(left)}{top}:{get_column_letter(right)}{max(bottom, target)}'
            if target > bottom:
                pattern = rb'(<dimension\b[^>]*\bref=")[^"]*(")'
                updated, count = re.subn(pattern, lambda m: m[1] + new_ref.encode() + m[2], updated)
                if count != 1:
                    raise RepairError('ambiguous_dimension_xml')
        result = BytesIO()
        with ZipFile(result, 'w') as output:
            output.comment = source.comment
            for info in source.infolist():
                output.writestr(copy(info), updated if info.filename == member else source.read(info.filename))
        candidate = result.getvalue()
        with ZipFile(BytesIO(candidate)) as check:
            if check.testzip() is not None or any(check.read(n) != source.read(n)
                                                  for n in source.namelist() if n != member):
                raise RepairError('archive_preservation_failed')
        # Verify all unrelated cells (including formulas) and exact requested values.
        before = load_workbook(BytesIO(original), read_only=True, data_only=False)
        after = load_workbook(BytesIO(candidate), read_only=True, data_only=False)
        try:
            for row in before['TRAIN_27']:
                for cell in row:
                    if getattr(cell, 'coordinate', None) not in changes and cell.value is not None:
                        other = after['TRAIN_27'][cell.coordinate]
                        if (other.value, other.data_type) != (cell.value, cell.data_type):
                            raise RepairError('unrelated_cell_changed')
            for reference, value in changes.items():
                if after['TRAIN_27'][reference].value != value:
                    raise RepairError('repair_verification_failed')
        finally:
            before.close()
            after.close()
        if plan(candidate)[1]:
            raise RepairError('idempotence_verification_failed')
        return candidate


def repair(path: Path) -> dict:
    path = Path(path)
    if path.is_symlink() or not path.is_file():
        raise RepairError('expected_regular_workbook')
    original = path.read_bytes()
    target, changes, existed = plan(original)
    if not changes:
        return {'changed': False, 'x2_valid': True, 'backup_created': False}
    candidate = patch_archive(original, target, changes)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ') + '-' + uuid.uuid4().hex[:8]
    backup = path.with_name(f'{path.stem}.backup-{stamp}{path.suffix}')
    stage = path.with_name(f'.{path.stem}.repair-{stamp}{path.suffix}')
    mode = stat.S_IMODE(path.stat().st_mode)
    for name, contents in ((backup, original), (stage, candidate)):
        with name.open('xb') as handle:
            os.chmod(name, mode)
            handle.write(contents)
            handle.flush()
            os.fsync(handle.fileno())
    if backup.read_bytes() != original or path.read_bytes() != original:
        raise RepairError('source_changed_or_backup_invalid')
    os.replace(stage, path)
    return {'changed': True, 'x2_valid': True, 'existing_row_updated': existed,
            'backup_created': True, 'backup_verified': True, 'other_parts_preserved': True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--private-root', default='private_data')
    parser.add_argument('--excel', default='private_data/ForteVision_TRAIN_27_detailed.xlsx')
    args = parser.parse_args()
    try:
        result = repair(private_path(args.excel, args.private_root))
        print(json.dumps(result))
    except RepairError as error:
        print(json.dumps({'repair_ok': False, 'error_code': str(error)}))
        raise SystemExit(2) from None
    except Exception as error:
        print(json.dumps({'repair_ok': False, 'error_code': type(error).__name__}))
        raise SystemExit(2) from None


if __name__ == '__main__':
    main()
