"""Validate local XLSX labels and MOV inventory; never infer missing labels."""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import time, timedelta
import math
from numbers import Real
from pathlib import Path
import re
import unicodedata
from typing import Any

from ml.local_io import private_path, run_cli, write_json
from ml.preprocess import probe_video


def normalize_video_id(value: Any) -> str:
    if value is None:
        raise ValueError("Empty identifier")
    if isinstance(value, Real) and not isinstance(value, bool):
        if not math.isfinite(value):
            raise ValueError("Nonfinite identifier")
        value = str(int(value)) if float(value).is_integer() else str(value)
    text = unicodedata.normalize("NFKC", str(value)).strip().upper()
    text = re.sub(r"\.MOV$", "", text).replace("Х", "X")
    text = re.sub(r"\s+", "", text)
    if not text or any(c in text for c in ("/", "\\", "\x00")) or text in {".", ".."}:
        raise ValueError("Invalid identifier")
    return text


def parse_timestamp(value: Any, numeric_unit: str = "seconds") -> float:
    if numeric_unit not in {"seconds", "excel_days"}:
        raise ValueError("Unknown numeric timestamp unit")
    if isinstance(value, time):
        result = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    elif isinstance(value, timedelta):
        result = value.total_seconds()
    elif isinstance(value, Real) and not isinstance(value, bool):
        result = float(value) * (86400 if numeric_unit == "excel_days" else 1)
    elif isinstance(value, str):
        fields = value.strip().replace(",", ".").split(":")
        if not 1 <= len(fields) <= 3:
            raise ValueError("Expected seconds, MM:SS or HH:MM:SS")
        if len(fields) == 1:
            return parse_timestamp(float(fields[0]), numeric_unit)
        numbers = [float(part) for part in fields]
        if any(not math.isfinite(n) or n < 0 for n in numbers):
            raise ValueError("Invalid timestamp")
        if any(n >= 60 for n in numbers[1:]) or any(not n.is_integer() for n in numbers[:-1]):
            raise ValueError("Invalid clock timestamp")
        result = sum(n * 60 ** i for i, n in enumerate(reversed(numbers)))
    else:
        raise ValueError("Unsupported timestamp")
    if not math.isfinite(result) or result < 0:
        raise ValueError("Timestamp must be nonnegative and finite")
    return result


def parse_timestamps(value: Any, numeric_unit: str) -> list[float]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return []
    parts = value.split(";") if isinstance(value, str) else [value]
    return sorted(set(parse_timestamp(part, numeric_unit) for part in parts))


def parse_label(value: Any, positive: set[str], negative: set[str]) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, Real) and not isinstance(value, bool):
        key = str(int(value)) if math.isfinite(value) and float(value).is_integer() else str(value)
    else:
        key = str(value).strip().casefold()
    if key in positive:
        return 1
    if key in negative:
        return 0
    raise ValueError("Unknown binary label")


def validate_dataset(
    excel: Path, videos_dir: Path, *, id_column: str, label_column: str,
    timestamp_column: str, sheet: str | None = None, header_row: int = 1,
    group_column: str | None = None, numeric_unit: str = "seconds",
    positive_values: tuple[str, ...] = ("1",), negative_values: tuple[str, ...] = ("0",),
) -> tuple[dict, dict]:
    from openpyxl import load_workbook

    positive = {s.strip().casefold() for s in positive_values}
    negative = {s.strip().casefold() for s in negative_values}
    if positive & negative or not positive or not negative or header_row < 1:
        raise ValueError("Invalid label mapping or header row")
    if not videos_dir.is_dir() or excel.suffix.lower() != ".xlsx":
        raise ValueError("Expected XLSX and existing video directory")
    issues: list[dict] = []
    labels: dict[str, list[dict]] = defaultdict(list)
    wb = load_workbook(excel, read_only=True, data_only=False)
    row_count = 0
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(min_row=header_row, values_only=True)
        headers = next(rows, ())
        wanted = [id_column, label_column, timestamp_column] + ([group_column] if group_column else [])
        if any(headers.count(name) != 1 for name in wanted):
            raise ValueError("Each configured column must occur exactly once")
        positions = {name: headers.index(name) for name in wanted}
        for row_number, row in enumerate(rows, header_row + 1):
            values = {name: row[index] for name, index in positions.items()}
            if all(v is None or v == "" for v in values.values()):
                continue
            row_count += 1
            try:
                identifier = normalize_video_id(values[id_column])
            except ValueError:
                issues.append({"code": "invalid_identifier", "row": row_number})
                continue
            entry: dict = {"row": row_number, "label": None, "timestamps_sec": [], "group_id": None, "valid": True}
            try:
                if any(isinstance(v, str) and v.startswith("=") for v in values.values()):
                    raise ValueError("Formula cells require explicit exported values")
                entry["label"] = parse_label(values[label_column], positive, negative)
                entry["timestamps_sec"] = parse_timestamps(values[timestamp_column], numeric_unit)
                if group_column and values[group_column] is not None:
                    entry["group_id"] = str(values[group_column]).strip() or None
                if entry["label"] == 1 and not entry["timestamps_sec"]:
                    raise ValueError("Positive label requires a timestamp")
                if entry["label"] == 0 and entry["timestamps_sec"]:
                    raise ValueError("Negative label has event timestamps")
            except (ValueError, TypeError):
                entry["valid"] = False
                issues.append({"code": "invalid_label_or_timestamp", "row": row_number, "video_id": identifier})
            labels[identifier].append(entry)
    finally:
        wb.close()

    files: dict[str, list[Path]] = defaultdict(list)
    for path in sorted(videos_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() == ".mov":
            if not path.resolve().is_relative_to(videos_dir.resolve()):
                issues.append({"code": "video_symlink_outside_directory"})
                continue
            try:
                identifier = normalize_video_id(path.stem)
            except ValueError:
                issues.append({"code": "invalid_video_identifier"})
                continue
            files[identifier].append(path)
    if not files:
        issues.append({"code": "no_videos"})
    if not labels:
        issues.append({"code": "no_labels"})
    records = []
    for identifier in sorted(files.keys() | labels.keys()):
        matched_files, matched_labels = files.get(identifier, []), labels.get(identifier, [])
        if len(matched_files) > 1:
            issues.append({"code": "duplicate_video_id", "video_id": identifier, "paths": [str(p) for p in matched_files]})
        if len(matched_labels) > 1:
            issues.append({"code": "duplicate_label_id", "video_id": identifier, "rows": [e["row"] for e in matched_labels]})
        if not matched_files:
            issues.append({"code": "missing_video", "video_id": identifier})
        entry = matched_labels[0] if len(matched_labels) == 1 else None
        if not entry or entry["label"] is None:
            issues.append({"code": "missing_label", "video_id": identifier})
        metadata = None
        if len(matched_files) == 1:
            try:
                metadata = probe_video(matched_files[0])
            except (ValueError, OSError, RuntimeError):
                issues.append({"code": "unreadable_video", "video_id": identifier})
        eligible = bool(entry and entry["valid"] and entry["label"] is not None and metadata)
        if entry and metadata and any(t >= metadata["duration_sec"] for t in entry["timestamps_sec"]):
            issues.append({"code": "timestamp_out_of_range", "video_id": identifier})
            eligible = False
        records.append({
            "video_id": identifier, "video_path": str(matched_files[0].resolve()) if len(matched_files) == 1 else None,
            "label": entry["label"] if entry else None,
            "timestamps_sec": entry["timestamps_sec"] if entry else [],
            "group_id": entry["group_id"] if entry else None,
            "eligible": eligible, "metadata": metadata,
        })
    valid_entries = [e for group in labels.values() for e in group if e["valid"] and e["label"] is not None]
    summary = {"excel_rows": row_count, "unique_label_ids": len(labels),
               "video_files": sum(map(len, files.values())), "labeled_rows": len(valid_entries),
               "positive_rows": sum(e["label"] == 1 for e in valid_entries),
               "negative_rows": sum(e["label"] == 0 for e in valid_entries),
               "eligible_videos": sum(r["eligible"] for r in records), "issue_count": len(issues)}
    report = {"schema_version": 1, "summary": summary, "issues": issues}
    manifest = {"schema_version": 1, "validation_ok": not issues,
                "issue_codes": sorted({issue["code"] for issue in issues}),
                "numeric_timestamp_unit": numeric_unit, "videos": records}
    return report, manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--private-root", default="private_data")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--videos", required=True)
    parser.add_argument("--output", required=True, help="New directory for local report and manifest")
    for name in ("id", "label", "timestamp"):
        parser.add_argument(f"--{name}-column", required=True)
    parser.add_argument("--group-column")
    parser.add_argument("--sheet")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--numeric-timestamp-unit", choices=["seconds", "excel_days"], default="seconds")
    parser.add_argument("--positive-value", action="append")
    parser.add_argument("--negative-value", action="append")
    args = parser.parse_args()
    excel, videos, output = [private_path(p, args.private_root) for p in (args.excel, args.videos, args.output)]
    if output.exists():
        raise FileExistsError("Output already exists")
    report, manifest = validate_dataset(excel, videos, id_column=args.id_column,
        label_column=args.label_column, timestamp_column=args.timestamp_column,
        group_column=args.group_column, sheet=args.sheet, header_row=args.header_row,
        numeric_unit=args.numeric_timestamp_unit, positive_values=tuple(args.positive_value or ["1"]),
        negative_values=tuple(args.negative_value or ["0"]))
    output.mkdir(parents=True)
    write_json(output / "report.json", report)
    write_json(output / "manifest.json", manifest)
    print("Local validation report saved. " + ("Issues require local review." if report["issues"] else "Validation passed."))
    return 2 if report["issues"] else 0


if __name__ == "__main__":
    run_cli(main)
